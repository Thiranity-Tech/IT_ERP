from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"K:\IT\IT_ERP\docs\Frontend_Feature_Handoff.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        el = tc_mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def style_runs(paragraph, size=10.5, bold=False, color=None):
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)
for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5) if style_name != "Heading 3" else RGBColor(0x1F, 0x4D, 0x78)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
run = title.add_run("IT ERP Frontend Handoff")
run.font.name = "Calibri"
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_before = Pt(0)
subtitle.paragraph_format.space_after = Pt(12)
subrun = subtitle.add_run(
    "This document explains, page by page, what the frontend developer should build for the IT ERP app."
)
subrun.font.name = "Calibri"
subrun.font.size = Pt(10.5)
subrun.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

def add_section(title_text, intro_text, bullets):
    doc.add_paragraph(title_text, style="Heading 1")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run("Build it like this: ").bold = True
    p.add_run(intro_text)
    for item in bullets:
        doc.add_paragraph(item, style="List Bullet")


add_section(
    "1. Login page",
    "a simple company login screen with one username field, one password field, a Sign in button, and a Forgot password link.",
    [
        "The username placeholder should say Enter Username or Employee ID.",
        "The password placeholder should say Enter password.",
        "Keep the page clean, centered, and professional.",
        "Do not add extra text, duplicate help text, or clutter under the form.",
    ],
)

add_section(
    "2. Forgot password page",
    "a reset screen with Employee ID, email-based code delivery, a Send code button, and a Back to sign in link.",
    [
        "Use email only for the reset flow for now.",
        "Show error messages clearly and keep the page responsive after an error.",
        "After the code is sent, move the user to the next reset step without making them refresh.",
        "Do not show the removed mobile/SMS path on this page.",
    ],
)

add_section(
    "3. Dashboard page",
    "one main dashboard area with summary cards and a short overview, not repeated labels.",
    [
        "Show the ERP dashboard title and the workspace overview.",
        "Keep the top summary cards short and useful.",
        "Remove labels that do not help the user, like Live from backend.",
        "Do not repeat the same summary sentence in multiple places.",
    ],
)

add_section(
    "4. Users page",
    "a clean list of users with the important employee information visible when it exists.",
    [
        "Show employee ID, name, role, department, team, phone, email, and status if those values exist.",
        "Super admin should be able to create, edit, and view users.",
        "Keep the table or card layout simple and easy to scan.",
        "If a field is empty, hide it instead of showing a dash or empty placeholder.",
    ],
)

add_section(
    "5. Departments page",
    "a department list with clear edit and delete actions on each row or card.",
    [
        "Show the department name clearly.",
        "Rename should be possible from the edit action.",
        "Delete should be available where allowed.",
        "Keep the page simple and easy to understand.",
    ],
)

add_section(
    "6. Teams page",
    "a team list with the team name, its department, and edit/delete actions.",
    [
        "Show the team name and the department name together.",
        "Allow edit and delete actions where the role has permission.",
        "Do not freeze or lock the team list if the user is allowed to change it.",
    ],
)

add_section(
    "7. Attendance page",
    "attendance records plus the import and manual entry experience in one clear page.",
    [
        "Show Employee ID, name, date, late minutes, overtime minutes, and status.",
        "Manual entry should calculate late arrival and overtime from the office timing settings.",
        "Do not ask the user to type late minutes or overtime minutes by hand.",
        "Show only the useful inputs and hide fields that are not needed.",
    ],
)

add_section(
    "8. Attendance timing page",
    "the office in-time, out-time, and grace minutes settings.",
    [
        "Super admin should be able to view the timing settings.",
        "Admin should be able to edit the timing settings.",
        "Other roles should not be able to change these values.",
    ],
)

add_section(
    "9. Payroll page",
    "a payroll workspace with salary settings, generated payroll, and the payslip layout.",
    [
        "Show the base salary settings and let the super admin manage them.",
        "Keep the base salary list empty at first, then add items after they are saved.",
        "Show the generated payroll list below the settings area.",
        "The generated payroll name should use Date + Employee ID.",
        "Show the payslip with employee name, employee ID, month, earnings, deductions, and final salary.",
        "Show overtime and loss of pay clearly in the payroll breakdown.",
    ],
)

add_section(
    "10. Projects page",
    "a project list with the main project details and the actions the user can take.",
    [
        "Show project name, client, department, status, deadline, and priority.",
        "Each project should have view, edit, and delete actions where allowed.",
        "Completed projects should still be visible in archive or history.",
    ],
)

add_section(
    "11. Tasks page",
    "a task list with task ID in ascending order and clear create/edit screens.",
    [
        "Show task title, project, assignee, status, and deadline.",
        "Sort task IDs in ascending order.",
        "Create task and edit task pages must have proper page headers.",
        "Make it obvious who created the task and who completed it.",
    ],
)

add_section(
    "12. Support tickets page",
    "a ticket list that makes the CRM workflow easy to follow.",
    [
        "Show ticket title, description, status, and priority.",
        "CRM team users should see the create ticket option.",
        "Managers should be able to approve or move tickets forward in the workflow.",
        "Show in-progress and completed states clearly.",
    ],
)

add_section(
    "13. Employee lookup page",
    "a simple employee detail screen that only shows real data.",
    [
        "Show only the fields that actually have values.",
        "If phone number, email, or designation is empty, hide that row.",
        "The phone number should stay on one line if possible.",
        "Show contact details, role, team, and department when they exist.",
        "Do not show an empty records box on the website.",
    ],
)

add_section(
    "14. Notifications page",
    "a notification list where each item takes the user to the right page.",
    [
        "Clicking a notification should open the correct page.",
        "If the product flow removes viewed notifications, make that happen automatically.",
        "Keep the notification list short and easy to scan.",
    ],
)

add_section(
    "15. Settings page",
    "settings for changing password and other role-based options.",
    [
        "Show the password change option in settings.",
        "Only show settings in the roles that should actually use them.",
        "Do not show admin-only items in the wrong menu.",
    ],
)

add_section(
    "16. General rules for all pages",
    "a clean, stable UI that feels like one product instead of many mixed screens.",
    [
        "Use the same blue-green color palette across the app.",
        "Keep navigation stable so the sidebar does not jump back to the top on every click.",
        "Keep page headers on every form or detail screen.",
        "Do not leave blank boxes or repeated helper text on the screen.",
        "Hide empty values instead of showing fake placeholders or dashes.",
    ],
)

doc.save(OUTPUT)
print(OUTPUT)
